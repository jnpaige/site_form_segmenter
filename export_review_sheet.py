#!/usr/bin/env python3
"""
export_review_sheet.py — one spreadsheet carrying every pipeline output, so a
whole corpus can be read against the source PDFs in a single document.

Long form: one row per addressable item — paragraph, heading, table row, or
figure — with where it came from, what pass0 decided about its page, and what
stage2 tagged it with (resolved against stage1's assemblage record, so the
locality/layer/culture behind a tag are on the row rather than in a lookup).

Sheets:
    items        the long-form row-per-item sheet
    assemblages  one row per (document, assemblage) — stage1's full ROAD record
    documents    one row per document — page/table/figure retention

Usage:
    uv run --with openpyxl python export_review_sheet.py
    uv run --with openpyxl python export_review_sheet.py --corpus holdout --out review.xlsx
"""
import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / "lib"))
from page_parser import parse_pages
from block_parser import split_blocks, stitch_continuations

ROOT = Path(r"C:\Users\jpaige\Box\4. Research\LLM_segmenting_procedural_units")
CORPORA = {"pilot": ROOT / "pdfs_ocr_figures_pilot", "holdout": ROOT / "pdfs_ocr_figures_holdout"}

ITEM_COLUMNS = [
    # --- source ---
    "corpus", "document", "page", "item_id", "item_type",
    # --- pass 0 ---
    "pass0_section", "pass0_section_label", "page_kept_by_pass0",
    # --- stage 2 tags, resolved through stage 1 ---
    "tag_count", "tag_keys", "tag_labels",
    "tag_localities", "tag_arch_layers", "tag_geo_layers", "tag_cultures",
    "is_comparative", "is_general", "is_untagged",
    # --- figure-only ---
    "figure_file", "figure_classification", "figure_heading_path",
    # --- content ---
    "char_count", "text",
]

ASSEMBLAGE_COLUMNS = [
    "corpus", "document", "cluster_id", "label_as_reported", "also_known_as",
    "locality", "archaeological_layer", "geological_layer",
    "archaeological_stratigraphy", "n_artefacts", "evidence",
    "blocks_tagged",
]

DOCUMENT_COLUMNS = [
    "corpus", "document", "pages_total", "pages_kept", "pct_pages_kept",
    "blocks_total", "blocks_kept", "blocks_tagged", "pct_blocks_tagged_end_to_end",
    "table_rows_total", "table_rows_kept", "figures_total", "figures_on_kept_pages",
    "assemblages", "headings_recovered", "continuations_stitched",
]


# OCR text carries control characters where the source had symbols the decoder
# didn't map (a stray 0x0b where "±" was, most often). Excel refuses them
# outright, so they are stripped on the way into a cell rather than left to
# abort the export on one row in the middle of a corpus.
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(value):
    return _ILLEGAL.sub(" ", value) if isinstance(value, str) else value


def _join(values) -> str:
    return " | ".join(str(v) for v in values if v not in (None, ""))


def _page_sections(pass0: dict | None) -> dict[int, tuple[str, str]]:
    """{page -> (category, label)}. A page can carry more than one category;
    keep them all so the sheet shows what pass0 actually said."""
    out: dict[int, list[tuple[str, str]]] = {}
    for category, entries in (pass0 or {}).items():
        if not isinstance(entries, list):
            continue
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            for p in entry.get("pages", []):
                out.setdefault(int(p), []).append((category, entry.get("label") or ""))
    return {p: (_join(c for c, _ in v), _join(dict.fromkeys(l for _, l in v)))
            for p, v in out.items()}


def collect(corpus_name: str, corpus_dir: Path):
    items, assemblages, documents = [], [], []

    for doc in sorted(d for d in corpus_dir.iterdir() if d.is_dir() and not d.name[0].isdigit()):
        tags_path = next(doc.glob("*.cluster_tags.json"), None)
        if tags_path is None:
            continue
        data = json.loads(tags_path.read_text(encoding="utf-8"))

        clusters = {c["cluster_id"]: c for c in data.get("clusters", []) if c.get("cluster_id")}
        tags = {t["block_id"]: t for t in data.get("block_tags", []) if t.get("block_id")}
        sections = _page_sections(data.get("pass0_result"))
        kept_pages = {b["page"] for b in data.get("blocks", [])}
        tagged_per_cluster = {cid: 0 for cid in clusters}

        ocr_path = doc / "ocr_docling.json"
        ocr = json.loads(ocr_path.read_text(encoding="utf-8")) if ocr_path.exists() else []
        pages_total = len(ocr) or (max(kept_pages) + 1 if kept_pages else 0)
        table_rows = [(p["page_index"], len([r for r in t["rows"] if r.get("bbox")]))
                      for p in ocr for t in p.get("tables", [])]

        for block in data.get("blocks", []):
            tag = tags.get(block["block_id"], {})
            cids = [c for c in (tag.get("cluster_ids") or []) if c in clusters]
            for cid in cids:
                tagged_per_cluster[cid] += 1
            sec, label = sections.get(block["page"], ("", ""))
            picked = [clusters[c] for c in cids]
            items.append({
                "corpus": corpus_name, "document": doc.name, "page": block["page"],
                "item_id": block["block_id"], "item_type": block["type"],
                "pass0_section": sec, "pass0_section_label": label,
                "page_kept_by_pass0": block["page"] in kept_pages,
                "tag_count": len(cids),
                "tag_keys": _join(cids),
                "tag_labels": _join(c.get("label_as_reported") for c in picked),
                "tag_localities": _join(dict.fromkeys(
                    c.get("locality") for c in picked if c.get("locality"))),
                "tag_arch_layers": _join(c.get("archaeological_layer") for c in picked),
                "tag_geo_layers": _join(c.get("geological_layer") for c in picked),
                "tag_cultures": _join(dict.fromkeys(
                    c.get("archaeological_stratigraphy") for c in picked
                    if c.get("archaeological_stratigraphy"))),
                "is_comparative": len(cids) > 1,
                "is_general": bool(tag.get("general")),
                "is_untagged": not cids and not tag.get("general"),
                "figure_file": "", "figure_classification": "", "figure_heading_path": "",
                "char_count": len(block.get("text") or ""),
                "text": (block.get("text") or "")[:32000],
            })

        # Figures are not blocks and are not tagged by stage2; they are listed
        # so a reviewer can see what is on a page alongside its text.
        figures_path = doc / "figures" / "figures.json"
        figures = []
        if figures_path.exists():
            raw = json.loads(figures_path.read_text(encoding="utf-8"))
            figures = raw if isinstance(raw, list) else raw.get("figures", [])
        for i, fig in enumerate(figures):
            page = fig.get("page")
            sec, label = sections.get(page, ("", "")) if page is not None else ("", "")
            items.append({
                "corpus": corpus_name, "document": doc.name,
                "page": page if page is not None else "", "item_id": f"fig_{i:03d}",
                "item_type": "figure",
                "pass0_section": sec, "pass0_section_label": label,
                "page_kept_by_pass0": page in kept_pages if page is not None else "",
                "tag_count": "", "tag_keys": "", "tag_labels": "",
                "tag_localities": "", "tag_arch_layers": "", "tag_geo_layers": "",
                "tag_cultures": "", "is_comparative": "", "is_general": "",
                "is_untagged": "",
                "figure_file": fig.get("file") or "",
                "figure_classification": fig.get("classification") or "",
                "figure_heading_path": _join(fig.get("heading_path") or []),
                "char_count": len(fig.get("caption") or ""),
                "text": (fig.get("caption") or "")[:32000],
            })

        for cid, c in clusters.items():
            assemblages.append({
                "corpus": corpus_name, "document": doc.name, "cluster_id": cid,
                "label_as_reported": c.get("label_as_reported"),
                "also_known_as": _join(c.get("also_known_as") or []),
                "locality": c.get("locality"),
                "archaeological_layer": c.get("archaeological_layer"),
                "geological_layer": c.get("geological_layer"),
                "archaeological_stratigraphy": c.get("archaeological_stratigraphy"),
                "n_artefacts": c.get("n_artefacts"),
                "evidence": _join(c.get("evidence") or []),
                "blocks_tagged": tagged_per_cluster.get(cid, 0),
            })

        blocks_total = 0
        md = next(doc.glob("*.md"), None)
        if md is not None:
            all_blocks, _ = stitch_continuations(split_blocks(parse_pages(md)))
            blocks_total = len(all_blocks)

        n_tagged = sum(1 for b in data.get("blocks", [])
                       if (tags.get(b["block_id"], {}).get("cluster_ids")
                           or tags.get(b["block_id"], {}).get("general")))
        fig_pages = [f.get("page") for f in figures if f.get("page") is not None]
        documents.append({
            "corpus": corpus_name, "document": doc.name,
            "pages_total": pages_total, "pages_kept": len(kept_pages),
            "pct_pages_kept": round(100 * len(kept_pages) / pages_total) if pages_total else "",
            "blocks_total": blocks_total, "blocks_kept": len(data.get("blocks", [])),
            "blocks_tagged": n_tagged,
            # Against EVERY block in the document, not just the ones pass0 kept —
            # a ratio over kept blocks flatters a document that dropped pages,
            # because the blocks it never saw are absent from the denominator.
            "pct_blocks_tagged_end_to_end": (round(100 * n_tagged / blocks_total)
                                             if blocks_total else ""),
            "table_rows_total": sum(n for _, n in table_rows),
            "table_rows_kept": sum(n for pg, n in table_rows if pg in kept_pages),
            "figures_total": len(fig_pages),
            "figures_on_kept_pages": sum(1 for p in fig_pages if p in kept_pages),
            "assemblages": len(clusters),
            "headings_recovered": data.get("n_headings_recovered"),
            "continuations_stitched": data.get("n_continuations_stitched"),
        })

    return items, assemblages, documents


def write_workbook(sheets: dict[str, tuple[list[str], list[dict]]], out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for name, (columns, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top")
        for row in rows:
            ws.append([_clean(row.get(c, "")) for c in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
        for i, col in enumerate(columns, start=1):
            width = 60 if col == "text" else min(max(len(col) + 2, 10), 28)
            ws.column_dimensions[get_column_letter(i)].width = width
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--corpus", choices=[*CORPORA, "both"], default="both")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    names = list(CORPORA) if args.corpus == "both" else [args.corpus]
    items, assemblages, documents = [], [], []
    for name in names:
        i, a, d = collect(name, CORPORA[name])
        items += i
        assemblages += a
        documents += d
        print(f"{name:8s} {len(d)} documents, {len(i)} items, {len(a)} assemblages")

    out = Path(args.out) if args.out else Path(f"pipeline_review_{args.corpus}.xlsx")
    write_workbook({
        "items": (ITEM_COLUMNS, items),
        "assemblages": (ASSEMBLAGE_COLUMNS, assemblages),
        "documents": (DOCUMENT_COLUMNS, documents),
    }, out)
    print(f"\nwrote {out.resolve()}")


if __name__ == "__main__":
    main()
